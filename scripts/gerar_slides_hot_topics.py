#!/usr/bin/env python3
"""Gera o deck de 4 slides de Hot Topics de Busca do ICP (bloco do business plan).

Fonte de conteúdo: docs/Slides_Hot_Topics_Busca_Business_Plan_Caveo.md
Identidade visual herdada de scripts/md_to_docx_caveo.py (verde 0E8A5F, Calibri).

Uso:  python3 scripts/gerar_slides_hot_topics.py
Saída: outputs/hot-topics-busca-20260807-caveo/caveo-hot-topics-busca-icp.pptx
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------- identidade
VERDE = RGBColor(0x0E, 0x8A, 0x5F)
VERDE_BG = RGBColor(0xE8, 0xF4, 0xEF)
ESCURO = RGBColor(0x1A, 0x1A, 0x1A)
CINZA = RGBColor(0x55, 0x55, 0x55)
CINZA_BG = RGBColor(0xF4, 0xF4, 0xF4)
CINZA_LINHA = RGBColor(0xD9, 0xD9, 0xD9)
VERMELHO = RGBColor(0xB0, 0x3A, 0x2B)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
FONTE = "Calibri"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


# ------------------------------------------------------------------ helpers
def txt(slide, x, y, w, h, text, *, size=14, bold=False, color=ESCURO,
        align=PP_ALIGN.LEFT, italic=False, spacing=1.0, anchor=MSO_ANCHOR.TOP):
    """Caixa de texto simples; `text` aceita lista de linhas."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        run = p.add_run()
        run.text = line
        run.font.name = FONTE
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill, *, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.text_frame.text = ""
    return s


def header(slide, numero, kicker, titulo, mensagem):
    """Faixa de topo padrão: número, kicker, título e a mensagem única."""
    rect(slide, Emu(0), Emu(0), SLIDE_W, Inches(0.09), VERDE)
    txt(slide, Inches(0.6), Inches(0.36), Inches(1.0), Inches(0.3),
        f"{numero:02d}", size=12, bold=True, color=VERDE)
    txt(slide, Inches(1.1), Inches(0.36), Inches(9.0), Inches(0.3),
        kicker.upper(), size=11, bold=True, color=CINZA)
    txt(slide, Inches(0.6), Inches(0.68), Inches(12.1), Inches(0.55),
        titulo, size=30, bold=True, color=ESCURO)
    txt(slide, Inches(0.6), Inches(1.32), Inches(12.1), Inches(0.34),
        mensagem, size=14, italic=True, color=VERDE)


def footer(slide, texto):
    rect(slide, Inches(0.6), Inches(6.93), Inches(12.13), Emu(9525), CINZA_LINHA)
    txt(slide, Inches(0.6), Inches(7.05), Inches(12.13), Inches(0.28),
        texto, size=9.5, color=CINZA)


def table(slide, x, y, w, headers, rows, col_ratios, *, size=11,
          header_size=10, row_h=0.295, header_h=0.32, zero_col=None,
          total_row=False, header_fill=VERDE, header_color=BRANCO):
    """Tabela achatada, sem estilo nativo do PowerPoint."""
    n_rows = len(rows) + 1
    shape = slide.shapes.add_table(n_rows, len(headers), x, y, w,
                                   Inches(header_h + row_h * len(rows)))
    tbl = shape.table
    tbl.first_row = False
    tbl.horz_banding = False

    total_ratio = sum(col_ratios)
    for i, ratio in enumerate(col_ratios):
        tbl.columns[i].width = Emu(int(w * ratio / total_ratio))
    tbl.rows[0].height = Inches(header_h)
    for r in range(1, n_rows):
        tbl.rows[r].height = Inches(row_h)

    def style(cell, text, *, bold, color, fill, sz, align):
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
        cell.margin_left = cell.margin_right = Inches(0.07)
        cell.margin_top = cell.margin_bottom = 0
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = str(text)
        run.font.name = FONTE
        run.font.size = Pt(sz)
        run.font.bold = bold
        run.font.color.rgb = color

    for c, head in enumerate(headers):
        style(tbl.cell(0, c), head, bold=True, color=header_color,
              fill=header_fill, sz=header_size,
              align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT)

    for r, row in enumerate(rows, start=1):
        is_total = total_row and r == n_rows - 1
        fill = VERDE_BG if is_total else (BRANCO if r % 2 else CINZA_BG)
        for c, val in enumerate(row):
            color = ESCURO
            bold = is_total or c == 0 and is_total
            if zero_col is not None and c == zero_col and str(val).strip() == "0":
                color, bold = VERMELHO, True
            style(tbl.cell(r, c), val, bold=bold or is_total, color=color,
                  fill=fill, sz=size,
                  align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT)
    return shape


def stat(slide, x, y, w, h, rotulo, valor, detalhe):
    rect(slide, x, y, w, h, CINZA_BG)
    rect(slide, x, y, Inches(0.045), h, VERDE)
    txt(slide, x + Inches(0.22), y + Inches(0.14), w - Inches(0.4), Inches(0.22),
        rotulo.upper(), size=9.5, bold=True, color=CINZA)
    txt(slide, x + Inches(0.22), y + Inches(0.38), w - Inches(0.4), Inches(0.34),
        valor, size=18, bold=True, color=ESCURO)
    txt(slide, x + Inches(0.22), y + Inches(0.76), w - Inches(0.4), Inches(0.22),
        detalhe, size=10, color=CINZA)


# ------------------------------------------------------------------- slides
def slide1(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, 1, "Demanda de busca do ICP médico",
           "Onde o médico realmente busca",
           "O médico não pesquisa funcionalidade. Ele pesquisa oferta — na proporção de 12 para 1.")

    # --- barras de proporção
    bar_x, bar_max = Inches(2.55), Inches(5.1)
    txt(s, Inches(0.6), Inches(2.15), Inches(1.9), Inches(0.3),
        "OFERTA", size=13, bold=True, color=ESCURO)
    rect(s, bar_x, Inches(2.12), bar_max, Inches(0.42), VERDE)
    txt(s, bar_x + bar_max + Inches(0.15), Inches(2.14), Inches(1.6), Inches(0.32),
        "5.460  ·  92%", size=13, bold=True, color=VERDE)

    txt(s, Inches(0.6), Inches(2.84), Inches(1.9), Inches(0.3),
        "FUNCIONALIDADE", size=13, bold=True, color=ESCURO)
    rect(s, bar_x, Inches(2.81), Inches(0.44), Inches(0.42), CINZA)
    txt(s, bar_x + Inches(0.59), Inches(2.83), Inches(1.8), Inches(0.32),
        "470  ·  8%", size=13, bold=True, color=CINZA)

    txt(s, Inches(0.6), Inches(3.42), Inches(7.2), Inches(0.3),
        "Buscas/mês no Brasil com marcador médico explícito  ·  total endereçável: 5.930/mês",
        size=10.5, color=CINZA)

    # --- três evidências
    txt(s, Inches(0.6), Inches(3.95), Inches(7.2), Inches(0.26),
        "TRÊS EVIDÊNCIAS INDEPENDENTES, MESMA DIREÇÃO", size=10, bold=True, color=VERDE)
    largura, gap = Inches(2.28), Inches(0.18)
    dados = [
        ("Volume de busca", "5.460 × 470", "12:1 a favor de oferta"),
        ("Termos que converteram", "100% × 0", "90 dias, não-marca"),
        ("CAC do Search de intenção", "R$ 536", "29% de fechamento"),
    ]
    for i, (rot, val, det) in enumerate(dados):
        stat(s, Inches(0.6) + i * (largura + gap), Inches(4.3),
             largura, Inches(1.08), rot, val, det)

    # --- painel de limpeza
    px, pw = Inches(8.35), Inches(4.38)
    rect(s, px, Inches(2.05), pw, Inches(3.33), None, line=CINZA_LINHA)
    txt(s, px + Inches(0.24), Inches(2.24), pw - Inches(0.48), Inches(0.28),
        "COMO O VOLUME FOI LIMPO", size=10, bold=True, color=VERDE)
    txt(s, px + Inches(0.24), Inches(2.52), pw - Inches(0.48), Inches(0.44),
        "De ~17.700 buscas/mês que contêm a palavra “médico”, 2/3 foram descartadas por não serem o ICP:",
        size=10, color=CINZA, spacing=1.1)
    table(s, px + Inches(0.24), Inches(3.08), pw - Inches(0.48),
          ["Descartado", "Vol/mês"],
          [["Salário e currículo de médico", "4.400"],
           ["Série de TV “Plantão Médico”", "3.750"],
           ["Paciente declarando IR", "1.800"],
           ["CNPJ de empresa de saúde", "880"],
           ["Convênio PJ, clínica, outros", "900"],
           ["Total descartado", "11.730"]],
          [3.1, 1.0], size=10, header_size=9.5, row_h=0.288, total_row=True)

    footer(s, "Fonte: Semrush base BR (82 keywords semente + varredura de cauda longa em 9 sementes) e termos de busca da conta Google Ads 3921127876, 09/05–06/08/2026.")
    return s


def slide2(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, 2, "Bloco Oferta — 5.460 buscas/mês",
           "Palavras-chave do bloco OFERTA",
           "O maior cluster e o cluster que mais converte são diferentes — e os dois são de oferta.")

    txt(s, Inches(0.6), Inches(1.86), Inches(6.0), Inches(0.26),
        "OS 5 CLUSTERS", size=10, bold=True, color=VERDE)
    table(s, Inches(0.6), Inches(2.16), Inches(6.0),
          ["Cluster", "Vol/mês", "CPC", "KD"],
          [["1. Contabilidade especializada em médicos", "2.780", "1,60–2,90", "19–26"],
           ["2. Abrir a PJ do médico", "1.630", "2,23–6,96", "17–21"],
           ["3. Tributos e enquadramento do médico", "850", "0,21–1,46", "36"],
           ["4. Documentos e regularidade", "140", "0,58", "14"],
           ["5. PJ ou CLT (objeção de entrada)", "60", "—", "—"],
           ["Total do bloco", "5.460", "", ""]],
          [3.5, 0.85, 1.0, 0.65], size=10.5, row_h=0.315, total_row=True)

    stat(s, Inches(0.6), Inches(4.36), Inches(2.9), Inches(1.08),
         "Maior volume", "Cluster 1", "2.780/mês · CPC R$ 1,60")
    stat(s, Inches(3.7), Inches(4.36), Inches(2.9), Inches(1.08),
         "Maior conversão", "Cluster 2", "CAC R$ 536 · travado por posição")

    txt(s, Inches(7.0), Inches(1.86), Inches(5.73), Inches(0.26),
        "10 TERMOS CONCENTRAM 71% DO BLOCO", size=10, bold=True, color=VERDE)
    table(s, Inches(7.0), Inches(2.16), Inches(5.73),
          ["Palavra-chave", "Vol/mês", "Cluster"],
          [["contabilidade medica", "1.000", "1"],
           ["contabilidade para medicos", "880", "1"],
           ["cnae medico", "390", "3"],
           ["contabilidade para medico", "320", "1"],
           ["cnpj medico", "320", "2"],
           ["cnpj para medicos", "320", "2"],
           ["contabilidade medico", "170", "1"],
           ["medico cooperado imposto de renda", "170", "3"],
           ["pj medico", "170", "2"],
           ["contabilidade medicos", "140", "1"],
           ["+ cauda de 46 termos (10–140/mês)", "1.580", "1–5"]],
          [3.6, 0.9, 0.7], size=10.5, row_h=0.315)

    txt(s, Inches(7.0), Inches(5.86), Inches(5.73), Inches(0.8),
        ["Cauda de intenção mais alta:  cnpj para medico 70 · pj para medicos 70 ·",
         "contador para medicos 70 · como abrir cnpj medico 50 · imposto pj medico 50 ·",
         "abrir pj medico 40 · pj medica 40 · cnpj para medico autonomo 40"],
        size=9.5, color=CINZA, spacing=1.15)

    footer(s, "Critério de inclusão: o termo precisa conter marcador médico (médico, médica, plantonista, cooperado) e expressar intenção de prestador de serviço.")
    return s


def slide3(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, 3, "Bloco Funcionalidade — 470 buscas/mês",
           "Palavras-chave do bloco FUNCIONALIDADE",
           "Funcionalidade não tem demanda de busca. Não é falta de volume no mercado — é ausência de pergunta.")

    txt(s, Inches(0.6), Inches(1.86), Inches(6.4), Inches(0.26),
        "O QUE EXISTE", size=10, bold=True, color=VERDE)
    table(s, Inches(0.6), Inches(2.16), Inches(6.4),
          ["Palavra-chave", "Vol/mês"],
          [["nota fiscal medico", "90"],
           ["medico pessoa fisica e obrigado a emitir nota fiscal", "40"],
           ["app para medicos", "40"],
           ["como emitir nota fiscal medico", "20"],
           ["como emitir nota fiscal de serviços medicos", "20"],
           ["emitir nota fiscal medico", "20"],
           ["como medico emite nota fiscal", "20"],
           ["medico é obrigado a dar nota fiscal", "20"],
           ["nota fiscal eletronica medicos", "20"],
           ["+ 9 outras variantes de nota fiscal médica", "180"],
           ["Total do bloco", "470"]],
          [4.9, 1.0], size=10.5, row_h=0.315, total_row=True)

    txt(s, Inches(7.4), Inches(1.86), Inches(5.33), Inches(0.26),
        "O QUE NÃO EXISTE — ZERO VOLUME NO ICP", size=10, bold=True, color=VERMELHO)
    table(s, Inches(7.4), Inches(2.16), Inches(5.33),
          ["Tema do produto", "Vol/mês"],
          [["Conta PJ para médico", "0"],
           ["Simulador / “quanto sobra pra você”", "0"],
           ["Gestão e controle financeiro", "0"],
           ["Recebimento garantido", "0"],
           ["Lembrete de pagamentos", "0"],
           ["Faturamento consolidado", "0"]],
          [4.0, 0.9], size=10.5, row_h=0.315, zero_col=1,
          header_fill=VERMELHO)

    cx, cy, cw = Inches(7.4), Inches(4.32), Inches(5.33)
    rect(s, cx, cy, cw, Inches(1.62), CINZA_BG)
    rect(s, cx, cy, Inches(0.045), Inches(1.62), VERMELHO)
    txt(s, cx + Inches(0.24), cy + Inches(0.18), cw - Inches(0.48), Inches(0.28),
        "A ARMADILHA DO VOLUME GENÉRICO", size=10, bold=True, color=VERMELHO)
    txt(s, cx + Inches(0.24), cy + Inches(0.5), cw - Inches(0.48), Inches(1.0),
        ["“emitir nota fiscal mei” tem 165.000 buscas/mês.",
         "“simulador simples nacional”, 3.600. “abrir conta pj”, 2.900.",
         "Todos sem marcador médico — e médico não pode ser MEI.",
         "Perseguir esse volume é comprar o público errado."],
        size=11, color=ESCURO, spacing=1.18)

    txt(s, Inches(0.6), Inches(6.06), Inches(6.4), Inches(0.6),
        ["Dos 470/mês, ~60 são de obrigatoriedade (“médico é obrigado a dar nota fiscal”) —",
         "argumento de Oferta, não demo de produto. Reclassificando, a razão vai de 12:1 para 13:1."],
        size=9.5, color=CINZA, spacing=1.15)

    footer(s, "Consequência de alocação: funcionalidade não vai para PMax. Vai para vídeo no Meta e para onboarding — onde a demanda é criada, não colhida.")
    return s


def slide4(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, 4, "Consequência para o plano",
           "O mercado é grande, mas está latente",
           "Busca não escala o plano — ela é o canal mais eficiente, não o maior.")

    txt(s, Inches(0.6), Inches(1.86), Inches(6.0), Inches(0.26),
        "A DEMANDA DECLARADA É ~2% DO SAM", size=10, bold=True, color=VERDE)
    table(s, Inches(0.6), Inches(2.16), Inches(6.0),
          ["Camada", "Tamanho"],
          [["TAM — médicos ativos no Brasil", "635.700"],
           ["SAM — médicos que fazem plantão", "275.900"],
           ["Demanda declarada em busca", "5.930/mês"],
           ["Equivalente do SAM por mês", "~2%"]],
          [4.3, 1.3], size=11.5, row_h=0.35, total_row=True)

    txt(s, Inches(0.6), Inches(3.78), Inches(6.0), Inches(0.44),
        ["Buscas são consultas, não pessoas — o número de médicos distintos",
         "em mercado por mês é ainda menor. É um teto, não uma estimativa."],
        size=9.5, italic=True, color=CINZA, spacing=1.15)

    bx, by, bw = Inches(0.6), Inches(4.4), Inches(6.0)
    rect(s, bx, by, bw, Inches(1.06), VERDE_BG)
    rect(s, bx, by, Inches(0.045), Inches(1.06), VERDE)
    txt(s, bx + Inches(0.24), by + Inches(0.16), bw - Inches(0.48), Inches(0.24),
        "TETO DO LADO BUSCA — taxa real do cnpj_medico", size=10, bold=True, color=VERDE)
    txt(s, bx + Inches(0.24), by + Inches(0.44), bw - Inches(0.48), Inches(0.5),
        ["5.930 buscas/mês × 5,3% clique→oportunidade",
         "= poucas centenas de oportunidades/mês"],
        size=11, bold=True, color=ESCURO, spacing=1.15)

    txt(s, Inches(0.6), Inches(5.7), Inches(6.0), Inches(0.6),
        ["Com capacidade de atendimento fora do gargalo, o teto do negócio é a eficiência de CAC",
         "no leilão — e o leilão de busca acaba. Metas acima de algumas centenas de clientes/mês",
         "dependem obrigatoriamente de criação de demanda."],
        size=10, color=CINZA, spacing=1.15)

    txt(s, Inches(7.0), Inches(1.86), Inches(5.73), Inches(0.26),
        "PRIORIDADE DE ORÇAMENTO EM GOOGLE", size=10, bold=True, color=VERDE)
    table(s, Inches(7.0), Inches(2.16), Inches(5.73),
          ["#  Onde", "Por quê"],
          [["1  Marca — fechar gap de impression share", "CAC R$ 141"],
           ["2  Search cnpj_medico — subir até travar", "CAC R$ 536 · 29%"],
           ["3  Novo Search: contabilidade especializada", "2.780/mês livre"],
           ["4  PMax — assets de Oferta, marca excluída", "teto por criativo"]],
          [3.75, 1.55], size=10.5, row_h=0.35)

    txt(s, Inches(7.0), Inches(3.78), Inches(5.73), Inches(0.26),
        "ALOCAÇÃO QUE OS DADOS SUSTENTAM", size=10, bold=True, color=VERDE)
    table(s, Inches(7.0), Inches(4.08), Inches(5.73),
          ["Pilar", "Canal", "Controle"],
          [["Oferta", "Search + PMax", "CAC/cliente"],
           ["Funcionalidade", "Meta vídeo + onboarding", "custo/SQL"]],
          [1.35, 2.6, 1.35], size=10.5, row_h=0.35)

    txt(s, Inches(7.0), Inches(5.42), Inches(5.73), Inches(1.0),
        ["Exclusão de marca no PMax é obrigatória: “caveo” tem 4.400 buscas/mês",
         "e a campanha institucional entrega cliente a R$ 141. Sem excluir, o PMax",
         "canibaliza esse resultado, infla o próprio número e o teste não",
         "responde nada."],
        size=10, color=CINZA, spacing=1.15)

    footer(s, "Conecta com as seções 2, 3 e 5 do business-plan-midia-paga-tam-sam-som.md. CAC blended do período: Meta R$ 5.411 · Google R$ 267 · geral R$ 974.")
    return s


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    for build in (slide1, slide2, slide3, slide4):
        build(prs)

    out = Path("outputs/hot-topics-busca-20260807-caveo")
    out.mkdir(parents=True, exist_ok=True)
    destino = out / "caveo-hot-topics-busca-icp.pptx"
    prs.save(destino)
    print(f"OK — {len(prs.slides.__iter__.__self__._sldIdLst)} slides em {destino}")


if __name__ == "__main__":
    main()
