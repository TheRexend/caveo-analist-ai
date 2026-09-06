#!/usr/bin/env python3
"""Deck de cartoes de criativo (formato visual escuro, thumbnail + metricas + texto).

Um slide por criativo: fundo escuro com identidade Caveo, thumbnail vertical
do anuncio com metricas do framework de teste ao lado, e narrativa (o que
testamos / o que funcionou (ou nao) / o que ensina sobre o medico PJ) num
painel a direita.
"""
import math
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONTE = "Calibri"

BG1 = RGBColor(0x0A, 0x1B, 0x36)
BG2 = RGBColor(0x11, 0x33, 0x5C)
VERDE = RGBColor(0x2E, 0xCC, 0x9A)
VERDE_ESCURO = RGBColor(0x0E, 0x8A, 0x5F)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
CINZA_CLARO = RGBColor(0xB8, 0xC6, 0xDC)
CINZA_MED = RGBColor(0x8A, 0x9B, 0xBA)
PAINEL = RGBColor(0x14, 0x2A, 0x4D)
AMARELO = RGBColor(0xF5, 0xA6, 0x23)
VERMELHO = RGBColor(0xE6, 0x5A, 0x4E)
CINZA_DOT = RGBColor(0x6B, 0x7B, 0x99)

COR_STATUS = {"verde": VERDE, "amarelo": AMARELO, "vermelho": VERMELHO, "na": CINZA_DOT}


def _wrap_lines(text, width_in, size_pt, chars_factor=0.50):
    """Estimativa de linhas apos quebra automatica (heuristica, sem metrica real de fonte)."""
    avg_char_in = (size_pt * chars_factor) / 72.0
    chars_per_line = max(8, int(width_in / avg_char_in))
    linhas = 0
    for paragrafo in text.split("\n"):
        linhas += max(1, math.ceil(len(paragrafo) / chars_per_line))
    return linhas


def _block_height(text, width_in, size_pt, spacing=1.22):
    linhas = _wrap_lines(text, width_in, size_pt)
    return Inches(linhas * (size_pt * spacing) / 72.0 + 0.06)


def _alpha(shape, pct):
    """Aplica transparencia (0-100, 100=opaco) ao preenchimento solido do shape."""
    sp = shape.fill.fore_color._xFill
    srgb = sp.find(qn('a:srgbClr'))
    if srgb is None:
        return
    for el in srgb.findall(qn('a:alpha')):
        srgb.remove(el)
    alpha_el = srgb.makeelement(qn('a:alpha'), {'val': str(int(pct * 1000))})
    srgb.append(alpha_el)


def _novo(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _rect(slide, x, y, w, h, fill, *, shape=MSO_SHAPE.RECTANGLE, alpha=None, line=None, radius=None):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            s.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
        if alpha is not None:
            _alpha(s, alpha)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.text_frame.text = ""
    return s


def _bg(slide):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    rect.shadow.inherit = False
    rect.line.fill.background()
    fill = rect.fill
    fill.gradient()
    stops = fill.gradient_stops
    stops[0].color.rgb = BG1
    stops[0].position = 0.0
    stops[1].color.rgb = BG2
    stops[1].position = 1.0
    fill.gradient_angle = 55
    for cx, cy, d, col in [
        (Inches(11.6), Inches(-1.4), Inches(4.2), BG2),
        (Inches(-1.6), Inches(5.7), Inches(3.6), BG2),
        (Inches(11.3), Inches(5.4), Inches(2.1), VERDE_ESCURO),
    ]:
        _rect(slide, cx, cy, d, d, col, shape=MSO_SHAPE.OVAL, alpha=16)
    return rect


def _txt(slide, x, y, w, h, text, *, size=14, bold=False, color=BRANCO,
          align=PP_ALIGN.LEFT, italic=False, spacing=1.08, anchor=MSO_ANCHOR.TOP,
          font=FONTE):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return box


def _header(slide, kicker, titulo, cor_kicker=VERDE):
    _rect(slide, Inches(0.6), Inches(0.42), Inches(0.42), Pt(2.4), cor_kicker)
    _txt(slide, Inches(1.14), Inches(0.33), Inches(9.0), Inches(0.3),
         kicker.upper(), size=12, bold=True, color=cor_kicker)
    n_linhas = _wrap_lines(titulo, 11.6, 27)
    _txt(slide, Inches(0.6), Inches(0.62), Inches(11.9), Inches(0.44 * n_linhas),
         titulo, size=27, bold=True, color=BRANCO, spacing=1.02)
    return 0.62 + 0.44 * n_linhas


def _thumb(slide, x, y, w, h, path, *, is_video):
    slide.shapes.add_picture(str(path), x, y, width=w, height=h)
    if is_video:
        d = Inches(0.6)
        cx = x + w / 2 - d / 2
        cy = y + h / 2 - d / 2
        circ = _rect(slide, cx, cy, d, d, BRANCO, shape=MSO_SHAPE.OVAL, alpha=88)
        circ.line.color.rgb = BRANCO
        circ.line.width = Pt(1.5)
        _txt(slide, cx, cy, d, d, "▶", size=19, bold=True,
             color=RGBColor(0x0A, 0x1B, 0x36), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def _metric_row(slide, x, y, w, label, valor, status):
    dot_d = Inches(0.14)
    _rect(slide, x, y + Inches(0.07), dot_d, dot_d, COR_STATUS[status], shape=MSO_SHAPE.OVAL)
    _txt(slide, x + Inches(0.26), y, w - Inches(0.26), Inches(0.22),
         label.upper(), size=10.5, bold=True, color=CINZA_CLARO)
    _txt(slide, x + Inches(0.26), y + Inches(0.21), w - Inches(0.26), Inches(0.36),
         valor, size=18, bold=True, color=BRANCO)


def slide_criativo(prs, item):
    s = _novo(prs)
    _bg(s)
    kicker = "O QUE FUNCIONOU" if item["bom"] else "O QUE NAO FUNCIONOU"
    cor_kicker = VERDE if item["bom"] else AMARELO
    fim_titulo = _header(s, kicker, item["titulo_criativo"], cor_kicker)
    sub_y = Inches(fim_titulo + 0.1)
    _txt(s, Inches(0.6), sub_y, Inches(11.0), Inches(0.3),
         f"{item['ad_name']} · entrou em {item['entrou']} · {item['campanha']}",
         size=12.5, color=CINZA_MED)

    top = Inches(fim_titulo + 0.5)

    # --- coluna 1: thumbnail ---
    tx, ty = Inches(0.6), top
    tw = Inches(2.15)
    th = Inches(2.15 * 16 / 9)
    _thumb(s, tx, ty, tw, th, item["thumb_path"], is_video=item["formato"] == "video")

    # --- coluna 2: metricas, ao lado do thumbnail ---
    mx = tx + tw + Inches(0.32)
    mw = Inches(1.75)
    metrics = item["metricas"]
    row_h = th / len(metrics)
    for i, (label, valor, status) in enumerate(metrics):
        _metric_row(s, mx, ty + row_h * i + Inches(0.12), mw, label, valor, status)

    # --- coluna 3: painel de texto ---
    px = mx + mw + Inches(0.35)
    py = top
    pw = SLIDE_W - px - Inches(0.6)
    ph = SLIDE_H - py - Inches(0.35)
    _rect(s, px, py, pw, ph, PAINEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, alpha=55, radius=0.04)

    inx, inw_in = px + Inches(0.38), (pw - Inches(0.76)) / Inches(1)
    y = py + Inches(0.3)

    def bloco(titulo, corpo, cor_titulo):
        nonlocal y
        _txt(s, inx, y, Inches(inw_in), Inches(0.3), titulo, size=15, bold=True, color=cor_titulo)
        alto = _block_height(corpo, inw_in, 13.5, spacing=1.3)
        _txt(s, inx, y + Inches(0.34), Inches(inw_in), alto, corpo, size=13.5, color=BRANCO, spacing=1.3)
        y += Inches(0.34) + alto + Inches(0.32)

    bloco("O que testamos", item["testamos"], VERDE)
    label_func = "O que funcionou" if item["bom"] else "O que não funcionou"
    bloco(label_func, item["funcionou"], AMARELO if not item["bom"] else VERDE)
    bloco("O que ensina sobre o médico PJ", item["ensina"], VERDE)

    return s


def slide_capa(prs, titulo, subtitulo, bullets):
    s = _novo(prs)
    _bg(s)
    _rect(s, Inches(0.6), Inches(2.55), Inches(0.55), Pt(3), VERDE)
    _txt(s, Inches(0.6), Inches(2.72), Inches(11.5), Inches(0.3),
         "RELATÓRIO SEMANAL DE CRIATIVOS · META ADS", size=13, bold=True, color=VERDE)
    _txt(s, Inches(0.6), Inches(3.02), Inches(11.8), Inches(1.1),
         titulo, size=38, bold=True, color=BRANCO)
    _txt(s, Inches(0.6), Inches(3.85), Inches(11.5), Inches(0.4),
         subtitulo, size=15, color=CINZA_CLARO)
    y = Inches(4.5)
    for b in bullets:
        _rect(s, Inches(0.6), y + Inches(0.09), Inches(0.09), Inches(0.09), VERDE, shape=MSO_SHAPE.OVAL)
        alto = _block_height(b, 10.6, 13, spacing=1.25)
        _txt(s, Inches(0.85), y, Inches(10.6), alto, b, size=13, color=BRANCO, spacing=1.25)
        y += alto + Inches(0.16)
    return s


def slide_notas(prs, titulo, blocos):
    """blocos: lista de (titulo, corpo, cor_titulo)."""
    s = _novo(prs)
    _bg(s)
    _header(s, "ressalvas", titulo, AMARELO)
    y = Inches(2.0)
    for tit, corpo, cor in blocos:
        _txt(s, Inches(0.6), y, Inches(11.9), Inches(0.28), tit, size=15, bold=True, color=cor)
        alto = _block_height(corpo, 11.9, 13, spacing=1.28)
        _txt(s, Inches(0.6), y + Inches(0.32), Inches(11.9), alto, corpo, size=13, color=BRANCO, spacing=1.28)
        y += Inches(0.32) + alto + Inches(0.28)
    return s


def gerar(itens, destino, *, capa=None, notas=None):
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    if capa:
        slide_capa(prs, *capa)
    for item in itens:
        slide_criativo(prs, item)
    if notas:
        slide_notas(prs, *notas)
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    prs.save(destino)
    return destino
