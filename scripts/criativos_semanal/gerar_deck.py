#!/usr/bin/env python3
"""Deck semanal de análise de criativo — 4 blocos + anexo.

Consome o dicionário montado pela skill criativos-semanal (Fase 5) e produz
outputs/criativos-semanal-AAAAMMDD-caveo/caveo-criativos-semanal.pptx
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from deck_caveo import (AMARELO, CINZA, ESCURO, SLIDE_H, SLIDE_W,  # noqa: E402
                        VERDE, VERMELHO, footer, header, table, txt)

COR_STATUS = {"verde": VERDE, "amarelo": AMARELO,
              "vermelho": VERMELHO, "na": CINZA}
MARCA = {"verde": "OK", "amarelo": "ATENCAO", "vermelho": "CRITICO", "na": "n/d"}
ROTULO_NIVEL = {"gancho": "Gancho — os 3 primeiros segundos",
                "retencao": "Retenção — o meio da peça",
                "interesse": "Interesse — a promessa",
                "conversao": "Conversão — criativo x LP"}


def _pct(v):
    return "—" if v is None else f"{v * 100:.1f}%".replace(".", ",")


def _brl(v):
    return "—" if v is None else f"R$ {v:,.2f}".replace(",", "@").replace(
        ".", ",").replace("@", ".")


def _valor(v, formato):
    if v is None:
        return "—"
    return _brl(v) if formato == "brl" else f"{v:,}".replace(",", ".")


def _novo(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def slide_consolidado(prs, d):
    s = _novo(prs)
    p = d["periodo"]
    header(s, 1, "consolidado do mês", "Onde estamos no mês",
           f"Mídia paga de {p['mes_inicio']} a {p['mes_fim']}.")
    linhas = [[c["metrica"], _valor(c["realizado"], c["formato"]),
               _valor(c["meta"], c["formato"]), _pct(c["pct_meta"]),
               _pct(c["ritmo"])] for c in d["consolidado"]]
    table(s, Inches(0.6), Inches(1.95), Inches(12.13),
          ["Métrica", "Realizado", "Meta", "% da meta", "Ritmo"],
          linhas, [3, 2, 2, 2, 2])
    txt(s, Inches(0.6), Inches(5.9), Inches(12.13), Inches(0.6),
        "Ritmo = realizado ÷ (meta × dias decorridos ÷ dias do mês). "
        "Ritmo de 100% é o passo exato para bater a meta até o fim do mês.",
        size=10, italic=True, color=CINZA)
    footer(s, "Definições de MQL e SQL: docs/fundacao-dados.md")
    return s


def _bloco_criativo(s, y, item, *, mostrar_causa):
    """Um card de criativo: nome, coorte, os 4 KPIs e a narrativa."""
    txt(s, Inches(0.6), y, Inches(6.0), Inches(0.3),
        f"{item['ad_name']} · entrou em {item['entrou']} "
        f"· coorte {item['coorte']}", size=13, bold=True, color=ESCURO)
    x = Inches(7.0)
    for kpi, rotulo in (("hook_rate", "Hook"), ("hold_rate_hook", "Hold"),
                        ("ctr_link", "CTR"), ("cpa", "CPA")):
        st = item["status"][kpi]
        valor = item["kpis"][kpi]
        texto = _brl(valor) if kpi == "cpa" else _pct(valor)
        txt(s, x, y, Inches(1.4), Inches(0.24), rotulo, size=9, color=CINZA)
        txt(s, x, y + Inches(0.2), Inches(1.4), Inches(0.28),
            f"{texto} · {MARCA[st]}", size=11, bold=True, color=COR_STATUS[st])
        x += Inches(1.42)
    corpo = y + Inches(0.34)
    if mostrar_causa:
        txt(s, Inches(0.6), corpo, Inches(6.2), Inches(0.9),
            f"Nível quebrado: {ROTULO_NIVEL[item['nivel_quebrado']]}",
            size=10, bold=True, color=VERMELHO)
        txt(s, Inches(0.6), corpo + Inches(0.28), Inches(12.13), Inches(0.9),
            item["causa"], size=11, color=ESCURO, spacing=1.15)
    else:
        txt(s, Inches(0.6), corpo, Inches(12.13), Inches(0.9),
            f"O que testamos: {item['testamos']}", size=11, color=ESCURO)
        txt(s, Inches(0.6), corpo + Inches(0.3), Inches(12.13), Inches(0.9),
            f"Por que funcionou: {item['por_que']}", size=11, color=ESCURO)
        txt(s, Inches(0.6), corpo + Inches(0.6), Inches(12.13), Inches(0.9),
            f"O que ensina sobre o médico PJ: {item['ensina']}",
            size=11, bold=True, color=VERDE, spacing=1.15)


def slide_funcionou(prs, d):
    s = _novo(prs)
    header(s, 2, "o que funcionou", "Os destaques da janela",
           "Destaque sai da coorte madura, que já passou a fase de aprendizado.")
    itens = d["funcionou"]
    if not itens:
        txt(s, Inches(0.6), Inches(2.4), Inches(12.13), Inches(0.6),
            "Nenhum criativo da janela atingiu volume e desempenho para virar "
            "destaque. Preferimos dizer isso a promover peça antiga.",
            size=13, italic=True, color=CINZA, spacing=1.2)
    y = Inches(2.05)
    for item in itens:
        _bloco_criativo(s, y, item, mostrar_causa=False)
        y += Inches(1.62)
    footer(s, "Piso de leitura: 500 impressões. Abaixo disso, o criativo vai para o anexo.")
    return s


def slide_nao_funcionou(prs, d):
    s = _novo(prs)
    header(s, 3, "o que não funcionou", "O que aprendemos com o que caiu",
           "O produto aqui é aprendizado sobre o médico PJ, não justificativa.")
    itens = d["nao_funcionou"]
    if not itens:
        txt(s, Inches(0.6), Inches(2.4), Inches(12.13), Inches(0.6),
            "Nenhum criativo com volume suficiente ficou abaixo do benchmark "
            "na janela.", size=13, italic=True, color=CINZA)
    y = Inches(2.05)
    for item in itens:
        _bloco_criativo(s, y, item, mostrar_causa=True)
        y += Inches(1.85)
    footer(s, "A causa vem da cascata: o primeiro nível vermelho é o diagnóstico.")
    return s


def slide_hipoteses(prs, d):
    s = _novo(prs)
    header(s, 4, "hipóteses", "O que testar na semana seguinte",
           "Cada hipótese nasce de um nível quebrado, não de brainstorm solto.")
    ant = d["hipoteses"]["anterior"]
    y = Inches(1.95)
    if ant:
        txt(s, Inches(0.6), y, Inches(12.13), Inches(0.28),
            f"A hipótese da semana passada ({ant['id']})",
            size=11, bold=True, color=CINZA)
        txt(s, Inches(0.6), y + Inches(0.26), Inches(12.13), Inches(0.3),
            ant["hipotese"], size=11, color=ESCURO)
        txt(s, Inches(0.6), y + Inches(0.54), Inches(12.13), Inches(0.3),
            f"O que aconteceu: {ant['aconteceu']}", size=11,
            bold=True, color=VERDE)
        y += Inches(1.05)
    linhas = [[h["id"], ROTULO_NIVEL[h["nivel"]].split(" — ")[0],
               h["framework"], h["celula"], h["hipotese"]]
              for h in d["hipoteses"]["novas"]]
    table(s, Inches(0.6), y, Inches(12.13),
          ["ID", "Nível", "Framework", "Célula PDA", "Hipótese"],
          linhas, [1, 1.4, 2, 2.4, 6])
    footer(s, "Célula PDA = Persona × Desire × Awareness ainda não testada (matriz.py).")
    return s


def slide_anexo(prs, d):
    s = _novo(prs)
    a = d["anexo"]
    header(s, 5, "anexo", "Ressalvas e leitura em andamento",
           "O que ainda não dá para afirmar, e por quê.")
    y = Inches(1.95)
    txt(s, Inches(0.6), y, Inches(12.13), Inches(0.28),
        "Ainda em leitura — abaixo de 500 impressões", size=11,
        bold=True, color=CINZA)
    em_leitura = a["em_leitura"]
    txt(s, Inches(0.6), y + Inches(0.26), Inches(12.13), Inches(0.4),
        ", ".join(f"{i['ad_name']} ({i['impressoes']} impr.)"
                  for i in em_leitura) or "nenhum",
        size=10, color=ESCURO)
    y += Inches(0.85)

    cob = a["cobertura_utmcon"]
    cor = VERDE if cob >= 0.8 else VERMELHO
    txt(s, Inches(0.6), y, Inches(12.13), Inches(0.28),
        "Cobertura de UtmCon__c", size=11, bold=True, color=CINZA)
    txt(s, Inches(0.6), y + Inches(0.26), Inches(12.13), Inches(0.4),
        f"{_pct(cob)} das oportunidades pagas têm criativo identificado."
        + ("" if cob >= 0.8 else " Abaixo de 80%: o ranking de funil é indicativo."),
        size=10, color=cor)
    y += Inches(0.85)

    txt(s, Inches(0.6), y, Inches(12.13), Inches(0.28),
        "Conjuntos com criativos concorrentes", size=11, bold=True, color=CINZA)
    conc = a["conjuntos_concorrentes"]
    txt(s, Inches(0.6), y + Inches(0.26), Inches(12.13), Inches(0.5),
        "; ".join(f"{c['adset']}: {', '.join(c['ads'])}" for c in conc)
        or "nenhum — cada conjunto tem um conceito distinto",
        size=10, color=ESCURO, spacing=1.15)
    y += Inches(0.95)

    txt(s, Inches(0.6), y, Inches(12.13), Inches(0.28),
        "Campanhas incluídas", size=11, bold=True, color=CINZA)
    txt(s, Inches(0.6), y + Inches(0.26), Inches(12.13), Inches(0.7),
        " · ".join(a["campanhas"]), size=9.5, color=CINZA, spacing=1.15)
    footer(s, "O Meta suprime entrega de peças repetitivas no mesmo conjunto (Andromeda).")
    return s


def gerar(dados, destino):
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    for build in (slide_consolidado, slide_funcionou, slide_nao_funcionou,
                  slide_hipoteses, slide_anexo):
        build(prs, dados)
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    prs.save(destino)
    return destino
