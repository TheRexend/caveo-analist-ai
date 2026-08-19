#!/usr/bin/env python3
"""Primitivas visuais dos decks da Caveo — verde 0E8A5F, Calibri, 16:9.

Extraído de scripts/gerar_slides_hot_topics.py para ser compartilhado pelos
geradores recorrentes. Aquele script mantém a própria cópia por ora: é um
deliverable já entregue e refatorá-lo agora só adicionaria risco.
"""
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------- identidade
VERDE = RGBColor(0x0E, 0x8A, 0x5F)
VERDE_BG = RGBColor(0xE8, 0xF4, 0xEF)
ESCURO = RGBColor(0x1A, 0x1A, 0x1A)
CINZA = RGBColor(0x55, 0x55, 0x55)
CINZA_BG = RGBColor(0xF4, 0xF4, 0xF4)
CINZA_LINHA = RGBColor(0xD9, 0xD9, 0xD9)
VERMELHO = RGBColor(0xB0, 0x3A, 0x2B)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
# Amarelo do status 🟡, para as tabelas de benchmark do deck semanal.
AMARELO = RGBColor(0xC8, 0x86, 0x00)
FONTE = "Calibri"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


# ------------------------------------------------------------------ helpers
def txt(slide, x, y, w, h, text, *, size=14, bold=False, color=ESCURO,
        align=PP_ALIGN.LEFT, italic=False, spacing=1.0, anchor=MSO_ANCHOR.TOP):
    """Caixa de texto simples; `text` aceita lista de linhas."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        run = p.add_run()
        run.text = line
        run.font.name = FONTE
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill, *, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.text_frame.text = ""
    return s


def header(slide, numero, kicker, titulo, mensagem):
    """Faixa de topo padrão: número, kicker, título e a mensagem única."""
    rect(slide, Emu(0), Emu(0), SLIDE_W, Inches(0.09), VERDE)
    txt(slide, Inches(0.6), Inches(0.36), Inches(1.0), Inches(0.3),
        f"{numero:02d}", size=12, bold=True, color=VERDE)
    txt(slide, Inches(1.1), Inches(0.36), Inches(9.0), Inches(0.3),
        kicker.upper(), size=11, bold=True, color=CINZA)
    txt(slide, Inches(0.6), Inches(0.68), Inches(12.1), Inches(0.55),
        titulo, size=30, bold=True, color=ESCURO)
    txt(slide, Inches(0.6), Inches(1.32), Inches(12.1), Inches(0.34),
        mensagem, size=14, italic=True, color=VERDE)


def footer(slide, texto):
    rect(slide, Inches(0.6), Inches(6.93), Inches(12.13), Emu(9525), CINZA_LINHA)
    txt(slide, Inches(0.6), Inches(7.05), Inches(12.13), Inches(0.28),
        texto, size=9.5, color=CINZA)


def table(slide, x, y, w, headers, rows, col_ratios, *, size=11,
          header_size=10, row_h=0.295, header_h=0.32, zero_col=None,
          total_row=False, header_fill=VERDE, header_color=BRANCO):
    """Tabela achatada, sem estilo nativo do PowerPoint."""
    n_rows = len(rows) + 1
    shape = slide.shapes.add_table(n_rows, len(headers), x, y, w,
                                   Inches(header_h + row_h * len(rows)))
    tbl = shape.table
    tbl.first_row = False
    tbl.horz_banding = False

    total_ratio = sum(col_ratios)
    for i, ratio in enumerate(col_ratios):
        tbl.columns[i].width = Emu(int(w * ratio / total_ratio))
    tbl.rows[0].height = Inches(header_h)
    for r in range(1, n_rows):
        tbl.rows[r].height = Inches(row_h)

    def style(cell, text, *, bold, color, fill, sz, align):
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
        cell.margin_left = cell.margin_right = Inches(0.07)
        cell.margin_top = cell.margin_bottom = 0
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = str(text)
        run.font.name = FONTE
        run.font.size = Pt(sz)
        run.font.bold = bold
        run.font.color.rgb = color

    for c, head in enumerate(headers):
        style(tbl.cell(0, c), head, bold=True, color=header_color,
              fill=header_fill, sz=header_size,
              align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT)

    for r, row in enumerate(rows, start=1):
        is_total = total_row and r == n_rows - 1
        fill = VERDE_BG if is_total else (BRANCO if r % 2 else CINZA_BG)
        for c, val in enumerate(row):
            color = ESCURO
            bold = is_total or c == 0 and is_total
            if zero_col is not None and c == zero_col and str(val).strip() == "0":
                color, bold = VERMELHO, True
            style(tbl.cell(r, c), val, bold=bold or is_total, color=color,
                  fill=fill, sz=size,
                  align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT)
    return shape


def stat(slide, x, y, w, h, rotulo, valor, detalhe):
    rect(slide, x, y, w, h, CINZA_BG)
    rect(slide, x, y, Inches(0.045), h, VERDE)
    txt(slide, x + Inches(0.22), y + Inches(0.14), w - Inches(0.4), Inches(0.22),
        rotulo.upper(), size=9.5, bold=True, color=CINZA)
    txt(slide, x + Inches(0.22), y + Inches(0.38), w - Inches(0.4), Inches(0.34),
        valor, size=18, bold=True, color=ESCURO)
    txt(slide, x + Inches(0.22), y + Inches(0.76), w - Inches(0.4), Inches(0.22),
        detalhe, size=10, color=CINZA)
