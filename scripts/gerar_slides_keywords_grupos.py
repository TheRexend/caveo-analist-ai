#!/usr/bin/env python3
"""Slides de keywords dos 2 grupos de maior potencial, no estilo do painel Semrush.

Grupos: (A) Contabilidade / Contador médico  ·  (B) Abertura de CNPJ / PJ médico
Escopo: só ICP médico individual. Clínica/consultório FORA (decisão 2026-08-07).

Volumes/KD/intenção: Semrush base `br`, **forma acentuada** (a forma canônica em
pt-BR; a sem acento é outro registro e infla o número).
CPC: a API devolve USD — convertido para BRL pelo fator 5,08, calibrado contra o
export do painel do cliente (1,60→8,13 · 1,76→8,94 · 0,87→4,42 · 1,50→7,62).

Uso:  python3 scripts/gerar_slides_keywords_grupos.py
Saída: outputs/hot-topics-busca-20260807-caveo/caveo-keywords-grupos-potencial.pptx
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

USD_BRL = 5.08

# --------------------------------------------------------------- paleta
AZUL_CLARO = RGBColor(0x1B, 0x6E, 0x9C)   # canto claro do gradiente
AZUL_ESCURO = RGBColor(0x0A, 0x22, 0x3C)  # canto escuro
MARCA_DAGUA = RGBColor(0x14, 0x3F, 0x63)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
KICKER = RGBColor(0xA8, 0xC4, 0xDA)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
TH_BG = RGBColor(0xFA, 0xFA, 0xFB)
TH_BG_SORT = RGBColor(0xEC, 0xEE, 0xF1)
TH_TXT = RGBColor(0x6B, 0x72, 0x80)
LINHA = RGBColor(0xE5, 0xE7, 0xEB)
KW_AZUL = RGBColor(0x2B, 0x5C, 0xE0)
TXT = RGBColor(0x37, 0x41, 0x51)
CINZA_ND = RGBColor(0x9C, 0xA3, 0xAF)
VERDE_KD = RGBColor(0x10, 0xB9, 0x81)
VERDE_KD_FORTE = RGBColor(0x05, 0x96, 0x69)
LARANJA_KD = RGBColor(0xF5, 0x9E, 0x0B)
BADGE_C = RGBColor(0xFD, 0xE0, 0x8A)
BADGE_I = RGBColor(0xC7, 0xD2, 0xFE)
PRIO_ALTA = RGBColor(0x05, 0x96, 0x69)
PRIO_MEDIA = RGBColor(0x92, 0x71, 0x0A)
FONTE = "Calibri"

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)


# --------------------------------------------------------------- helpers XML
def spacing(run, centesimos):
    """Espaçamento entre letras (a:rPr@spc, em 1/100 pt)."""
    run.font._rPr.set("spc", str(centesimos))


def highlight(run, rgb):
    """Fundo colorido atrás do glifo — usado como 'badge' de intenção."""
    rPr = run.font._rPr
    hl = rPr.makeelement(qn("a:highlight"), {})
    clr = hl.makeelement(qn("a:srgbClr"), {"val": f"{rgb}"})
    hl.append(clr)
    rPr.append(hl)


def borda_inferior(cell, rgb=LINHA, pt=0.75):
    """a:lnB precisa vir ANTES do preenchimento na ordem do schema."""
    tcPr = cell._tc.get_or_add_tcPr()
    antiga = tcPr.find(qn("a:lnB"))
    if antiga is not None:
        tcPr.remove(antiga)
    ln = tcPr.makeelement(qn("a:lnB"), {"w": str(int(pt * 12700)), "cap": "flat",
                                        "cmpd": "sng", "algn": "ctr"})
    fill = ln.makeelement(qn("a:solidFill"), {})
    clr = fill.makeelement(qn("a:srgbClr"), {"val": f"{rgb}"})
    fill.append(clr)
    ln.append(fill)
    tcPr.insert(0, ln)


# --------------------------------------------------------------- helpers pptx
def fundo(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), SLIDE_W, SLIDE_H)
    bg.shadow.inherit = False
    bg.line.fill.background()
    f = bg.fill
    f.gradient()
    f.gradient_angle = 25.0
    stops = f.gradient_stops
    stops[0].color.rgb = AZUL_CLARO
    stops[0].position = 0.0
    stops[1].color.rgb = AZUL_ESCURO
    stops[1].position = 1.0
    if len(stops) > 2:
        stops[2].color.rgb = AZUL_ESCURO
        stops[2].position = 1.0
    bg.text_frame.text = ""

    # marca d'água discreta no canto direito
    mk = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.4), Inches(-1.5),
                                Inches(5.2), Inches(5.2))
    mk.shadow.inherit = False
    mk.line.fill.background()
    mk.fill.solid()
    mk.fill.fore_color.rgb = MARCA_DAGUA
    mk.text_frame.text = ""


def titulo(slide, kicker, texto):
    box = slide.shapes.add_textbox(Inches(0.62), Inches(0.4), Inches(11.5), Inches(0.3))
    tf = box.text_frame
    tf.margin_left = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = kicker.upper()
    r.font.name, r.font.size = FONTE, Pt(13)
    r.font.color.rgb = KICKER
    spacing(r, 280)

    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.74), Inches(12.2), Inches(0.8))
    tf = box.text_frame
    tf.margin_left = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = texto
    r.font.name, r.font.size, r.font.bold = FONTE, Pt(38), True
    r.font.color.rgb = BRANCO


def rodape(slide, texto):
    box = slide.shapes.add_textbox(Inches(0.62), Inches(7.0), Inches(12.1), Inches(0.32))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = texto
    r.font.name, r.font.size = FONTE, Pt(9.5)
    r.font.color.rgb = KICKER


def kd_cor(kd):
    if kd is None:
        return CINZA_ND
    return VERDE_KD_FORTE if kd < 15 else (VERDE_KD if kd < 30 else LARANJA_KD)


def tabela(slide, rows):
    """Card branco + tabela no estilo do painel."""
    cx, cy, cw, ch = Inches(0.5), Inches(1.93), Inches(12.33), Inches(5.0)
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, cw, ch)
    card.shadow.inherit = False
    card.adjustments[0] = 0.022
    card.line.fill.background()
    card.fill.solid()
    card.fill.fore_color.rgb = CARD
    card.text_frame.text = ""

    pad = Inches(0.18)
    tx, ty, tw = cx + pad, cy + pad, cw - pad * 2
    header_h, row_h = 0.42, 0.35
    cols = ["Palavra-chave", "Intenção", "Volume", "KD %", "CPC (BRL)", "Prioridade"]
    ratios = [5.40, 1.25, 1.35, 1.25, 1.42, 1.30]

    shape = slide.shapes.add_table(len(rows) + 1, len(cols), tx, ty, tw,
                                   Inches(header_h + row_h * len(rows)))
    tbl = shape.table
    tbl.first_row = False
    tbl.horz_banding = False
    total = sum(ratios)
    for i, r in enumerate(ratios):
        tbl.columns[i].width = Emu(int(tw * r / total))
    tbl.rows[0].height = Inches(header_h)
    for i in range(1, len(rows) + 1):
        tbl.rows[i].height = Inches(row_h)

    def base(cell, fill):
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
        cell.margin_left = cell.margin_right = Inches(0.09)
        cell.margin_top = cell.margin_bottom = 0
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.text_frame.word_wrap = False

    # ---- cabeçalho
    for c, nome in enumerate(cols):
        cell = tbl.cell(0, c)
        base(cell, TH_BG_SORT if nome == "Volume" else TH_BG)
        borda_inferior(cell, LINHA, 1.0)
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = (nome + " ⌄") if nome == "Volume" else nome
        r.font.name, r.font.size = FONTE, Pt(10.5)
        r.font.color.rgb = TH_TXT

    # ---- corpo
    for i, row in enumerate(rows, start=1):
        kw, intents, vol, kd, cpc, prio = row
        for c in range(len(cols)):
            base(tbl.cell(i, c), CARD)
            borda_inferior(tbl.cell(i, c))

        # palavra-chave (azul, como link do painel)
        p = tbl.cell(i, 0).text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = kw
        r.font.name, r.font.size = FONTE, Pt(11.5)
        r.font.color.rgb = KW_AZUL

        # intenção — badges com fundo colorido
        p = tbl.cell(i, 1).text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        if intents:
            for j, letra in enumerate(intents):
                if j:
                    sep = p.add_run()
                    sep.text = " "
                    sep.font.name, sep.font.size = FONTE, Pt(11.5)
                r = p.add_run()
                r.text = f" {letra} "
                r.font.name, r.font.size, r.font.bold = FONTE, Pt(10.5), True
                r.font.color.rgb = RGBColor(0x44, 0x40, 0x22) if letra == "C" \
                    else RGBColor(0x27, 0x33, 0x7A)
                highlight(r, BADGE_C if letra == "C" else BADGE_I)
        else:
            r = p.add_run()
            r.text = "n.d."
            r.font.name, r.font.size = FONTE, Pt(11)
            r.font.color.rgb = CINZA_ND

        # volume
        p = tbl.cell(i, 2).text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = f"{vol:,}".replace(",", ".")
        r.font.name, r.font.size = FONTE, Pt(11.5)
        r.font.color.rgb = TXT

        # KD + bolinha de cor
        p = tbl.cell(i, 3).text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = (f"{kd} " if kd is not None else "n.d. ")
        r.font.name, r.font.size = FONTE, Pt(11.5)
        r.font.color.rgb = TXT if kd is not None else CINZA_ND
        d = p.add_run()
        d.text = "●"
        d.font.name, d.font.size = FONTE, Pt(11)
        d.font.color.rgb = kd_cor(kd)

        # CPC em BRL
        p = tbl.cell(i, 4).text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = f"{cpc:.2f}".replace(".", ",")
        r.font.name, r.font.size = FONTE, Pt(11.5)
        r.font.color.rgb = TXT if cpc else CINZA_ND

        # prioridade
        p = tbl.cell(i, 5).text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = prio
        r.font.name, r.font.size, r.font.bold = FONTE, Pt(11), True
        r.font.color.rgb = PRIO_ALTA if prio == "Alta" else PRIO_MEDIA


# --------------------------------------------------------------- dados
# (palavra-chave, intenções, volume, KD, CPC USD, prioridade)
GRUPO_A = [
    ("contabilidade para médicos",                          "C",  720, 25, 1.60, "Alta"),
    ("contabilidade médica",                                "C",  480, 26, 1.76, "Alta"),
    ("contabilidade para médico",                           "I",  390, 22, 1.76, "Alta"),
    ("contabilidade médicos",                               "C",  210, 23, 1.60, "Alta"),
    ("contabilidade online para médicos",                   "",    70, None, 0.92, "Alta"),
    ("contador para médico",                                "",    70, None, 2.90, "Alta"),
    ("contabilidade digital para médicos",                  "",    50, None, 0.00, "Alta"),
    ("contabilidade especializada para médicos",            "",    50, None, 1.56, "Alta"),
    ("contabilidade para médicos pj",                       "",    50, None, 0.00, "Alta"),
    ("contabilidade para médicos e profissionais da saúde", "",    50, None, 1.00, "Média"),
    ("contador para médicos",                               "",    40, None, 2.90, "Alta"),
    ("contabilidade médico",                                "",    30, None, 1.85, "Média"),
]

GRUPO_B = [
    ("cnpj médico",                                   "I",  170, 17, 2.82, "Alta"),
    ("cnpj para médicos",                             "I",  110, 16, 2.23, "Alta"),
    ("para ser médico plantonista precisa ter cnpj",  "I",   90,  3, 0.00, "Alta"),
    ("qual o melhor cnpj para médicos",              "IC",   90, 12, 1.12, "Alta"),
    ("cnpj para médico",                              "",    70, None, 2.98, "Alta"),
    ("abrir cnpj médico",                             "",    50, None, 5.73, "Alta"),
    ("cnpj médico individual",                        "",    50, None, 1.41, "Alta"),
    ("cnpj para médico plantonista",                  "",    50, None, 1.91, "Alta"),
    ("pj médico",                                     "",    50, None, 4.62, "Alta"),
    ("imposto médico pj",                             "",    50, None, 0.83, "Média"),
    ("abertura de pj médico",                         "",    30, None, 0.00, "Alta"),
    ("abrir pj médico",                               "",    30, None, 0.00, "Alta"),
]


def brl(rows):
    return [(k, i, v, kd, round(c * USD_BRL, 2), p) for k, i, v, kd, c, p in rows]


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H

    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s)
    titulo(s, "Cenário atual", "KeyWords — Contabilidade / Contador")
    tabela(s, brl(GRUPO_A))
    rodape(s, "Grupo com 18 termos ICP e 2.340 buscas/mês  ·  os 12 acima somam 2.210 (94%)  ·  "
              "clínica e consultório excluídos do escopo, junto de 910 buscas/mês  ·  "
              "Semrush base BR, forma acentuada  ·  CPC convertido de USD a R$ 5,08")

    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s)
    titulo(s, "Cenário atual", "KeyWords — Abertura de CNPJ / PJ")
    tabela(s, brl(GRUPO_B))
    rodape(s, "Grupo com 26 termos ICP e 1.150 buscas/mês  ·  os 12 acima somam 840 (73%)  ·  "
              "excluídas 660 buscas/mês de consulta de CNPJ de operadora (Amil, Hapvida) "
              "e 60 de convênio médico PJ  ·  Semrush base BR, forma acentuada")

    out = Path("outputs/hot-topics-busca-20260807-caveo")
    out.mkdir(parents=True, exist_ok=True)
    destino = out / "caveo-keywords-grupos-potencial.pptx"
    prs.save(destino)
    print(f"OK — 2 slides em {destino}")


if __name__ == "__main__":
    main()
